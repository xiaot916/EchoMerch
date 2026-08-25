from __future__ import annotations

import csv
import io
import sqlite3
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from app.core.local_database import BUSINESS_DAY, STORE_ID, LocalDatabase, q


@dataclass(frozen=True)
class StoreDatasetDefinition:
    key: str
    label: str
    table: str
    description: str


DATASET_DEFINITIONS: tuple[StoreDatasetDefinition, ...] = (
    StoreDatasetDefinition("overview", "店铺经营数据", "store_daily_overviews", "店铺每日经营核心指标"),
    StoreDatasetDefinition("products", "商品销售数据", "store_daily_product_rankings", "商品销售、访客和转化表现"),
    StoreDatasetDefinition("traffic", "流量来源数据", "store_daily_traffic_sources", "各级流量来源及成交表现"),
    StoreDatasetDefinition("promotion_campaigns", "推广计划数据", "store_daily_promotion_campaigns", "推广计划的展现、点击、消耗和成交"),
    StoreDatasetDefinition("promotion_items", "推广商品数据", "store_daily_promotion_items", "推广商品的投放和成交表现"),
    StoreDatasetDefinition("customers", "客户分析数据", "store_daily_customer_overviews", "新老客和客户价值指标"),
    StoreDatasetDefinition("members", "会员分析数据", "store_daily_member_analysis_overviews", "会员规模、成交和复购指标"),
    StoreDatasetDefinition("member_channels", "会员渠道数据", "store_daily_member_channels", "会员招募和渠道表现"),
    StoreDatasetDefinition("customer_service", "客服数据", "store_daily_customer_service_overviews", "客服接待、转化和服务指标"),
    StoreDatasetDefinition("live", "直播数据", "store_daily_live_overviews", "直播间流量和成交指标"),
    StoreDatasetDefinition("content", "内容经营数据", "store_daily_content_overviews", "内容曝光、互动和种草成交"),
    StoreDatasetDefinition("shopping_gold", "购物金数据", "store_daily_shopping_gold_overviews", "购物金充值和成交指标"),
    StoreDatasetDefinition("bybt", "百亿补贴数据", "store_daily_bybt_overviews", "百亿补贴商品和成交指标"),
    StoreDatasetDefinition("bybt_items", "百亿补贴商品明细", "store_daily_bybt_items", "百亿补贴每日商品成交明细"),
    StoreDatasetDefinition("cps", "CPS 数据", "store_daily_cps_overviews", "CPS 引流、付款和结算指标"),
    StoreDatasetDefinition("flash_sale", "淘宝秒杀数据", "store_daily_taobao_flash_sale_overviews", "秒杀活动商品和成交指标"),
    StoreDatasetDefinition("flash_sale_items", "淘宝秒杀商品明细", "store_daily_taobao_flash_sale_items", "淘宝秒杀每日活动商品明细"),
    StoreDatasetDefinition("taobao_current_prices", "淘宝当前价快照", "store_daily_taobao_current_price_items", "淘宝商品当前价和预测价格区间"),
    StoreDatasetDefinition("taobao_risk_prices", "淘宝红线价风险", "store_daily_taobao_risk_price_items", "淘宝商品红线价风险与风险说明"),
    StoreDatasetDefinition("taobao_activity_snapshots", "淘宝活动商品快照", "store_daily_taobao_activity_item_snapshots", "百补/秒杀在线与待上线商品快照"),
    StoreDatasetDefinition("utry_sample", "U先派样数据", "store_daily_utry_sample_overviews", "U先派样商品与新客指标"),
    StoreDatasetDefinition("utry_repurchase", "U先复购数据", "store_daily_utry_repurchase_overviews", "U先商品复购指标"),
    StoreDatasetDefinition("flow_overview", "流量旧概览", "store_daily_flow_overviews", "旧版流量概览指标；与流量来源明细分开"),
)

_DEFINITIONS_BY_KEY = {item.key: item for item in DATASET_DEFINITIONS}
_SENSITIVE_MARKERS = ("cookie", "token", "csrf", "密码", "签名", "请求头", "原始", "raw")
MAX_PAGE_SIZE = 100
MAX_EXPORT_ROWS = 100_000


class StoreDatasetMissing(ValueError):
    pass


class StoreDataLimitExceeded(ValueError):
    pass


class StoreDataServiceError(RuntimeError):
    pass


class StoreDataService:
    def __init__(self, database_path: Path) -> None:
        self._database = LocalDatabase(database_path)

    def get_definition(self, dataset_key: str) -> StoreDatasetDefinition:
        definition = _DEFINITIONS_BY_KEY.get(dataset_key)
        if definition is None:
            raise StoreDatasetMissing(f"不支持的数据类型：{dataset_key}")
        return definition

    def list_datasets(self, store_id: int) -> list[dict[str, Any]]:
        with self._database.connect() as conn:
            result: list[dict[str, Any]] = []
            for definition in DATASET_DEFINITIONS:
                columns = self._safe_columns(conn, definition.table)
                if not self._is_queryable(columns):
                    continue
                bounds = conn.execute(
                    f'''select count(*) as row_count, min({q(BUSINESS_DAY)}) as first_date,
                        max({q(BUSINESS_DAY)}) as latest_date
                        from "{definition.table}"
                        where {q(STORE_ID)} = ?''',
                    (store_id,),
                ).fetchone()
                result.append({
                    "key": definition.key,
                    "label": definition.label,
                    "description": definition.description,
                    "row_count": int(bounds["row_count"] or 0),
                    "first_date": bounds["first_date"],
                    "latest_date": bounds["latest_date"],
                    "status": "empty",
                    "lag_days": None,
                    "columns": self._public_columns(columns),
                })
            reference_date = max(
                (date.fromisoformat(str(item["latest_date"])) for item in result if item["latest_date"]),
                default=None,
            )
            for item in result:
                if not item["row_count"] or not item["latest_date"]:
                    continue
                latest_date = date.fromisoformat(str(item["latest_date"]))
                lag_days = (reference_date - latest_date).days if reference_date else 0
                item["lag_days"] = lag_days
                item["status"] = "current" if lag_days == 0 else "stale"
            return result

    def preview(
        self,
        *,
        store_id: int,
        dataset_key: str,
        start_date: date | None,
        end_date: date | None,
        page: int,
        page_size: int,
        search: str | None,
    ) -> dict[str, Any]:
        definition = self.get_definition(dataset_key)
        page = max(page, 1)
        page_size = min(max(page_size, 1), MAX_PAGE_SIZE)
        with self._database.connect() as conn:
            columns = self._queryable_columns(conn, definition)
            start_date, end_date = self._normalize_range(conn, definition, store_id, start_date, end_date)
            where, params = self._where_clause(columns, store_id, start_date, end_date, search)
            total = int(conn.execute(
                f'''select count(*) from "{definition.table}" where {where}''', params
            ).fetchone()[0])
            selected = ", ".join(q(column) for column in columns)
            offset = (page - 1) * page_size
            rows = conn.execute(
                f'''select {selected} from "{definition.table}"
                    where {where}
                    order by {q(BUSINESS_DAY)} desc
                    limit ? offset ?''',
                (*params, page_size, offset),
            ).fetchall()
            return {
                "dataset": {
                    "key": definition.key,
                    "label": definition.label,
                    "description": definition.description,
                },
                "store_id": store_id,
                "start_date": start_date.isoformat() if start_date else None,
                "end_date": end_date.isoformat() if end_date else None,
                "page": page,
                "page_size": page_size,
                "total": total,
                "columns": self._public_columns(columns),
                "rows": [self._row_dict(columns, row) for row in rows],
            }

    def export(
        self,
        *,
        store_id: int,
        dataset_key: str,
        start_date: date | None,
        end_date: date | None,
        search: str | None,
        file_format: str,
    ) -> tuple[bytes, str, str]:
        definition = self.get_definition(dataset_key)
        with self._database.connect() as conn:
            columns = self._queryable_columns(conn, definition)
            start_date, end_date = self._normalize_range(conn, definition, store_id, start_date, end_date)
            where, params = self._where_clause(columns, store_id, start_date, end_date, search)
            total = int(conn.execute(
                f'''select count(*) from "{definition.table}" where {where}''', params
            ).fetchone()[0])
            if total > MAX_EXPORT_ROWS:
                raise StoreDataLimitExceeded(
                    f"筛选结果有 {total:,} 行，超过单次导出上限 {MAX_EXPORT_ROWS:,} 行，请缩小日期范围。"
                )
            selected = ", ".join(q(column) for column in columns)
            rows = conn.execute(
                f'''select {selected} from "{definition.table}"
                    where {where}
                    order by {q(BUSINESS_DAY)} desc''',
                params,
            ).fetchall()
            public_columns = self._public_columns(columns)
            values = [[self._public_value(row[column]) for column in columns] for row in rows]

        if file_format == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow([item["label"] for item in public_columns])
            writer.writerows(values)
            return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", ".csv"
        if file_format == "xlsx":
            return self._xlsx_bytes(definition, public_columns, values, start_date, end_date), (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ), ".xlsx"
        raise ValueError("导出格式只支持 csv 或 xlsx。")

    @staticmethod
    def _safe_columns(conn: sqlite3.Connection, table: str) -> list[str]:
        rows = conn.execute(f'pragma table_info("{table}")').fetchall()
        return [str(row[1]) for row in rows if not any(marker in str(row[1]).lower() for marker in _SENSITIVE_MARKERS)]

    def _queryable_columns(self, conn: sqlite3.Connection, definition: StoreDatasetDefinition) -> list[str]:
        columns = self._safe_columns(conn, definition.table)
        if not self._is_queryable(columns):
            raise StoreDatasetMissing(f"数据表暂不可预览：{definition.label}")
        return columns

    @staticmethod
    def _is_queryable(columns: list[str]) -> bool:
        return STORE_ID in columns and BUSINESS_DAY in columns and len(columns) > 2

    @staticmethod
    def _public_columns(columns: list[str]) -> list[dict[str, str]]:
        return [{"key": column, "label": column} for column in columns]

    @staticmethod
    def _normalize_range(
        conn: sqlite3.Connection,
        definition: StoreDatasetDefinition,
        store_id: int,
        start_date: date | None,
        end_date: date | None,
    ) -> tuple[date | None, date | None]:
        latest = conn.execute(
            f'''select max({q(BUSINESS_DAY)}) from "{definition.table}" where {q(STORE_ID)} = ?''',
            (store_id,),
        ).fetchone()[0]
        if latest is None:
            return start_date, end_date
        latest_day = date.fromisoformat(str(latest))
        return start_date or latest_day - timedelta(days=29), end_date or latest_day

    @staticmethod
    def _where_clause(
        columns: list[str],
        store_id: int,
        start_date: date | None,
        end_date: date | None,
        search: str | None,
    ) -> tuple[str, list[Any]]:
        clauses = [f"{q(STORE_ID)} = ?"]
        params: list[Any] = [store_id]
        if start_date:
            clauses.append(f"{q(BUSINESS_DAY)} >= ?")
            params.append(start_date.isoformat())
        if end_date:
            clauses.append(f"{q(BUSINESS_DAY)} <= ?")
            params.append(end_date.isoformat())
        if search and search.strip():
            searchable = [column for column in columns if column not in {STORE_ID, BUSINESS_DAY}]
            if searchable:
                clauses.append("(" + " or ".join(f"cast({q(column)} as text) like ?" for column in searchable) + ")")
                params.extend([f"%{search.strip()}%"] * len(searchable))
        return " and ".join(clauses), params

    @staticmethod
    def _public_value(value: Any) -> Any:
        if isinstance(value, bytes):
            return value.decode("utf-8", errors="replace")
        return value

    @classmethod
    def _row_dict(cls, columns: list[str], row: sqlite3.Row) -> dict[str, Any]:
        return {column: cls._public_value(row[column]) for column in columns}

    @staticmethod
    def _xlsx_bytes(
        definition: StoreDatasetDefinition,
        columns: list[dict[str, str]],
        values: list[list[Any]],
        start_date: date | None,
        end_date: date | None,
    ) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("Excel 导出依赖未安装，请安装 openpyxl。") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "数据明细"
        headers = [item["label"] for item in columns]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F6B5A")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{_excel_column(len(headers))}{max(len(values) + 1, 1)}"
        for row in values:
            sheet.append(row)
        for index, header in enumerate(headers, start=1):
            width = min(max(len(str(header)) + 2, 12), 28)
            sheet.column_dimensions[_excel_column(index)].width = width

        criteria = workbook.create_sheet("筛选条件")
        criteria.append(["数据类型", definition.label])
        criteria.append(["日期范围", f"{start_date or '-'} 至 {end_date or '-'}"])
        criteria.append(["导出行数", len(values)])
        criteria.column_dimensions["A"].width = 16
        criteria.column_dimensions["B"].width = 42
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()


def _excel_column(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result
