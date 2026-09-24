from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl


SOURCE_HEADERS = ("商品ID", "链接名称", "品类", "系列", "名称", "渠道")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--database-path", type=Path, required=True)
    parser.add_argument("--store-id", type=int, default=1)
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()

    source = args.input.resolve()
    database = args.database_path.resolve()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    headers = tuple(_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if headers[: len(SOURCE_HEADERS)] != SOURCE_HEADERS:
        raise ValueError(f"Unexpected headers: {headers}")

    rows: list[tuple[str, ...]] = []
    seen: set[str] = set()
    for source_row in sheet.iter_rows(min_row=2, values_only=True):
        values = tuple(_text(value) for value in source_row[: len(SOURCE_HEADERS)])
        if not values[0]:
            continue
        if values[0] in seen:
            raise ValueError(f"Duplicate product ID: {values[0]}")
        if any(not value for value in values):
            raise ValueError(f"Blank source field for product ID: {values[0]}")
        seen.add(values[0])
        rows.append(values)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = database.with_name(f"{database.stem}.before_product_update_{timestamp}{database.suffix}")
    shutil.copy2(database, backup)

    connection = sqlite3.connect(database)
    try:
        connection.execute("pragma foreign_keys = on")
        if connection.execute('select 1 from stores where [店铺ID] = ?', (args.store_id,)).fetchone() is None:
            raise ValueError(f"Store {args.store_id} does not exist")
        before = connection.execute(
            'select count(*) from store_product_catalog where [店铺ID] = ?', (args.store_id,)
        ).fetchone()[0]
        existing_ids = {
            row[0]
            for row in connection.execute(
                'select [商品ID] from store_product_catalog where [店铺ID] = ?', (args.store_id,)
            )
        }
        if args.replace:
            connection.execute(
                'delete from store_product_catalog where [店铺ID] = ?', (args.store_id,)
            )
        connection.executemany(
            '''
            insert into store_product_catalog
                ([店铺ID], [商品ID], [商品名称], [类型], [属性], [系列], [定位], [渠道])
            values (?, ?, ?, '', ?, ?, ?, ?)
            on conflict([店铺ID], [商品ID]) do update set
                [商品名称] = excluded.[商品名称],
                [属性] = excluded.[属性],
                [系列] = excluded.[系列],
                [定位] = excluded.[定位],
                [渠道] = excluded.[渠道]
            ''',
            [
                (args.store_id, product_id, link_name, name, series, category, channel)
                for product_id, link_name, category, series, name, channel in rows
            ],
        )
        connection.commit()
        after = connection.execute(
            'select count(*) from store_product_catalog where [店铺ID] = ?', (args.store_id,)
        ).fetchone()[0]
        matched = connection.execute(
            f'''select count(*) from store_product_catalog
                where [店铺ID] = ? and [商品ID] in ({','.join('?' for _ in rows)})''',
            (args.store_id, *(row[0] for row in rows)),
        ).fetchone()[0]
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

    print(json.dumps({
        "status": "updated",
        "source_rows": len(rows),
        "mode": "replace" if args.replace else "merge",
        "updated_existing": 0 if args.replace else len(existing_ids & seen),
        "inserted_new": len(rows) if args.replace else len(seen - existing_ids),
        "removed_not_in_source": len(existing_ids - seen) if args.replace else 0,
        "preserved_not_in_source": 0 if args.replace else len(existing_ids - seen),
        "rows_before": before,
        "rows_after": after,
        "source_ids_verified_in_database": matched,
        "database": str(database),
        "backup": str(backup),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
