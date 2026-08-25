from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "backend"))

from app.core.config import settings  # noqa: E402
from app.core.local_database import LocalDatabase, PRODUCT_CATALOG_TABLE, q  # noqa: E402


HEADERS = ("商品ID", "商品名称", "类型", "属性", "系列", "定位", "渠道")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main() -> int:
    parser = argparse.ArgumentParser(description="Import the store product catalog from an Excel workbook.")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--store-id", type=int, default=1)
    parser.add_argument("--database-path", type=Path, default=Path(settings.local_database_path))
    args = parser.parse_args()
    input_path = args.input.expanduser().resolve()
    database_path = args.database_path.expanduser().resolve()

    workbook = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header_values = tuple(_text(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1)))
    missing = [header for header in HEADERS if header not in header_values]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    indexes = [header_values.index(header) for header in HEADERS]

    rows: list[tuple[str, ...]] = []
    seen: set[str] = set()
    duplicate_count = 0
    skipped_count = 0
    for source_row in worksheet.iter_rows(min_row=2, values_only=True):
        values = tuple(_text(source_row[index]) if index < len(source_row) else "" for index in indexes)
        product_id = values[0]
        if not product_id:
            skipped_count += 1
            continue
        if product_id in seen:
            duplicate_count += 1
            continue
        seen.add(product_id)
        rows.append(values)

    database = LocalDatabase(database_path)
    database.initialize_schema()
    with database.connect(read_only=False) as conn:
        store_exists = conn.execute('select 1 from stores where "店铺ID" = ?', (args.store_id,)).fetchone()
        if store_exists is None:
            raise ValueError(f"Store {args.store_id} does not exist.")
        conn.execute(f'delete from {q(PRODUCT_CATALOG_TABLE)} where "店铺ID" = ?', (args.store_id,))
        conn.executemany(
            f'''insert into {q(PRODUCT_CATALOG_TABLE)}
                ("店铺ID", "商品ID", "商品名称", "类型", "属性", "系列", "定位", "渠道")
                values (?, ?, ?, ?, ?, ?, ?, ?)''',
            [(args.store_id, *row) for row in rows],
        )
        conn.commit()

    print(json.dumps({
        "status": "imported",
        "database": str(database_path),
        "store_id": args.store_id,
        "source": str(input_path),
        "rows": len(rows),
        "duplicates_skipped": duplicate_count,
        "blank_rows_skipped": skipped_count,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
